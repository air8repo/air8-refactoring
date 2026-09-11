"""额度管理 Excel 导出测试"""
import io
from openpyxl import load_workbook


SAMPLE_BUYER_ROWS = [
    {
        'buyer_code': 'B1', 'buyer_name': 'Amazon Services', 'currency': 'USD',
        'celling': 100000.0, 'reserved': 20000.0, 'actual': 75000.0,
        'total_occupied': 95000.0, 'headroom': 5000.0, 'occupancy_rate': 0.95,
        'pairs': [
            {'uid': 'U1', 'buyer_code': 'B1', 'buyer_name': 'Amazon Services',
             'supplier_code': 'S1', 'supplier_name': 'Photonverse Inc', 'currency': 'USD',
             'celling': 100000.0, 'reserved': 20000.0, 'actual': 75000.0,
             'total_occupied': 95000.0, 'headroom': 5000.0, 'occupancy_rate': 0.95},
        ],
    },
]


SAMPLE_PAIR_ROWS = [
    {'uid': 'U1', 'buyer_code': 'B1', 'buyer_name': 'Amazon Services', 'supplier_code': 'S1',
     'supplier_name': 'Photonverse Inc', 'currency': 'USD', 'celling': 100000.0, 'reserved': 20000.0,
     'actual': 75000.0, 'total_occupied': 95000.0, 'headroom': 5000.0, 'occupancy_rate': 0.95},
]


class TestBuildCreditLimitWorkbook:

    def test_creates_exactly_three_fixed_sheets(self):
        from backend.app.services.credit_export_service import build_credit_limit_workbook
        data = build_credit_limit_workbook(SAMPLE_BUYER_ROWS, SAMPLE_PAIR_ROWS, SAMPLE_DEAL_ROWS)
        wb = load_workbook(io.BytesIO(data))
        assert wb.sheetnames == ['Buyer', 'Buyer-Supplier', 'Deal']

    def test_buyer_sheet_contains_row_and_total(self):
        from backend.app.services.credit_export_service import build_credit_limit_workbook
        data = build_credit_limit_workbook(SAMPLE_BUYER_ROWS, SAMPLE_PAIR_ROWS, SAMPLE_DEAL_ROWS)
        wb = load_workbook(io.BytesIO(data))
        rows = list(wb['Buyer'].iter_rows(values_only=True))
        assert rows[0][0] == 'Buyer'
        assert rows[1][0] == 'Amazon Services'
        assert rows[-1][0] == 'Total'
        header = rows[0]
        celling_idx = header.index('Celling')
        assert rows[-1][celling_idx] == 100000.0

    def test_buyer_supplier_sheet_is_flat_not_split_by_buyer(self):
        from backend.app.services.credit_export_service import build_credit_limit_workbook
        data = build_credit_limit_workbook(SAMPLE_BUYER_ROWS, SAMPLE_PAIR_ROWS, SAMPLE_DEAL_ROWS)
        wb = load_workbook(io.BytesIO(data))
        rows = list(wb['Buyer-Supplier'].iter_rows(values_only=True))
        assert rows[0][0] == 'Buyer'
        header = rows[0]
        assert 'Supplier' in header
        assert rows[1][header.index('Supplier')] == 'Photonverse Inc'
        assert rows[-1][0] == 'Total'

    def test_deal_sheet_carries_buyer_supplier_identifiers_and_total(self):
        from backend.app.services.credit_export_service import build_credit_limit_workbook
        data = build_credit_limit_workbook(SAMPLE_BUYER_ROWS, SAMPLE_PAIR_ROWS, SAMPLE_DEAL_ROWS)
        wb = load_workbook(io.BytesIO(data))
        rows = list(wb['Deal'].iter_rows(values_only=True))
        header = rows[0]
        assert 'Buyer' in header
        assert 'Supplier' in header
        assert rows[1][header.index('FR#')] == 'FR1'
        assert rows[-1][0] == 'Total'

    def test_empty_input_still_produces_three_sheets_with_header_only(self):
        from backend.app.services.credit_export_service import build_credit_limit_workbook
        data = build_credit_limit_workbook([], [], [])
        wb = load_workbook(io.BytesIO(data))
        assert wb.sheetnames == ['Buyer', 'Buyer-Supplier', 'Deal']
        assert list(wb['Buyer'].iter_rows(values_only=True))[0][0] == 'Buyer'


SAMPLE_DEAL_ROWS = [
    {'finance_request_number': 'FR1', 'invoice_number': 'INV1', 'buyer_name': 'Buyer One',
     'supplier_name': 'Seller One', 'financing_amount': 1000.0, 'earmark_forecast': 1000.0,
     'credit_utilization': 0.0, 'to_be_settled_on_db': None, 'total_os': 1000.0,
     'status_display': 'eligible', 'finance_status_display': '', 'settlement_status': '',
     'due_date': '2026-10-01', 'actual_funding_date': '', 'batch_number': ''},
]


class TestBuildDealDetailWorkbook:

    def test_creates_deal_sheet_with_header_and_rows(self):
        from backend.app.services.credit_export_service import build_deal_detail_workbook
        from backend.app.services.credit_limit_service import compute_deal_totals
        total_row = compute_deal_totals(SAMPLE_DEAL_ROWS)
        data = build_deal_detail_workbook(SAMPLE_DEAL_ROWS, total_row)
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0][0] == 'FR#'
        assert rows[1][0] == 'FR1'

    def test_appends_total_row_at_the_end(self):
        from backend.app.services.credit_export_service import build_deal_detail_workbook
        from backend.app.services.credit_limit_service import compute_deal_totals
        total_row = compute_deal_totals(SAMPLE_DEAL_ROWS)
        data = build_deal_detail_workbook(SAMPLE_DEAL_ROWS, total_row)
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        last_row = rows[-1]
        assert last_row[0] == 'Total'
        header = rows[0]
        idx = header.index('Financing Amount')
        assert last_row[idx] == 1000.0
