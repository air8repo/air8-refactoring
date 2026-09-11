"""额度管理 Excel 导出 — Buyer/Buyer-Supplier/Deal 三个固定 sheet"""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from backend.app.services.credit_limit_service import compute_deal_totals

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
_THIN_BORDER = Border(left=Side(style='thin'), right=Side(style='thin'),
                       top=Side(style='thin'), bottom=Side(style='thin'))
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center")


def _write_header(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _THIN_BORDER
        cell.alignment = _CENTER


def _write_data_row(ws, values):
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.border = _THIN_BORDER
        cell.alignment = _LEFT


def _autofit_columns(ws):
    for column_cells in ws.columns:
        length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = length + 2


_BUYER_HEADERS = ['Buyer', 'Air8 Buyer ID', 'Currency', 'Celling', 'Reserved', 'Actual', 'Total Occupied', 'Headroom', 'Occupancy Rate']
_PAIR_HEADERS = ['Buyer', 'Air8 Buyer ID', 'Supplier', 'Air8 Seller ID', 'Currency', 'Celling', 'Reserved', 'Actual', 'Total Occupied', 'Headroom', 'Occupancy Rate']


def _amount_row_values(name, id_value, row):
    occupancy = row.get('occupancy_rate')
    return [
        name, id_value, row.get('currency', 'USD'),
        row.get('celling', 0.0), row.get('reserved', 0.0), row.get('actual', 0.0),
        row.get('total_occupied', 0.0),
        row.get('headroom') if row.get('headroom') is not None else 'N/A',
        f"{occupancy:.1%}" if occupancy is not None else 'N/A',
    ]


def _amount_total_row_values(rows, blank_count):
    """blank_count = 该 sheet 中 Celling 列之前除 'Total' 占用的那一格外，还有几个非数值列
    （名称/ID/币种等）需要留空对齐。Buyer sheet 是 Buyer/Air8 Buyer ID/Currency 共 3 列，
    'Total' 占 1 格，blank_count=2；Buyer-Supplier sheet 是 Buyer/Air8 Buyer
    ID/Supplier/Air8 Seller ID/Currency 共 5 列，blank_count=4。"""
    celling_sum = sum(r.get('celling', 0.0) for r in rows)
    reserved_sum = sum(r.get('reserved', 0.0) for r in rows)
    actual_sum = sum(r.get('actual', 0.0) for r in rows)
    total_sum = sum(r.get('total_occupied', 0.0) for r in rows)
    headroom_sum = round(celling_sum - total_sum, 2) if celling_sum else 'N/A'
    occupancy_avg = round(total_sum / celling_sum, 4) if celling_sum else None
    prefix = ['Total'] + [''] * blank_count
    return prefix + [
        round(celling_sum, 2), round(reserved_sum, 2), round(actual_sum, 2), round(total_sum, 2),
        headroom_sum, f"{occupancy_avg:.1%}" if occupancy_avg is not None else 'N/A',
    ]


def build_credit_limit_workbook(buyer_rows, pair_rows, deal_rows):
    """构建额度管理 Excel：固定 3 个 sheet（Buyer/Buyer-Supplier/Deal），各自末尾追加 Total 行。"""
    wb = Workbook()

    buyer_ws = wb.active
    buyer_ws.title = 'Buyer'
    _write_header(buyer_ws, _BUYER_HEADERS)
    for b in buyer_rows:
        _write_data_row(buyer_ws, _amount_row_values(b.get('buyer_name', ''), b.get('buyer_code', ''), b))
    if buyer_rows:
        _write_data_row(buyer_ws, _amount_total_row_values(buyer_rows, blank_count=2))
    _autofit_columns(buyer_ws)

    pair_ws = wb.create_sheet('Buyer-Supplier')
    _write_header(pair_ws, _PAIR_HEADERS)
    for p in pair_rows:
        occupancy = p.get('occupancy_rate')
        pair_ws.append([
            p.get('buyer_name', ''), p.get('buyer_code', ''), p.get('supplier_name', ''), p.get('supplier_code', ''),
            p.get('currency', 'USD'), p.get('celling', 0.0), p.get('reserved', 0.0), p.get('actual', 0.0),
            p.get('total_occupied', 0.0),
            p.get('headroom') if p.get('headroom') is not None else 'N/A',
            f"{occupancy:.1%}" if occupancy is not None else 'N/A',
        ])
        for cell in pair_ws[pair_ws.max_row]:
            cell.border = _THIN_BORDER
            cell.alignment = _LEFT
    if pair_rows:
        _write_data_row(pair_ws, _amount_total_row_values(pair_rows, blank_count=4))
    _autofit_columns(pair_ws)

    deal_ws = wb.create_sheet('Deal')
    _write_header(deal_ws, _DEAL_HEADERS)
    for d in deal_rows:
        _write_data_row(deal_ws, _deal_row_values(d))
    if deal_rows:
        _write_data_row(deal_ws, _deal_total_row_values(compute_deal_totals(deal_rows)))
    _autofit_columns(deal_ws)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


_DEAL_HEADERS = ['FR#', 'Invoice#', 'Buyer', 'Supplier', 'Financing Amount', 'Earmark Forecast',
                  'Credit Utilization', 'To Be Settled on DB', 'Total O/S', 'Status',
                  'Finance Details - Finance Status', 'Invoice Details - Settlement Status',
                  'Due Date', 'Funding Date', 'Batch']


def _deal_row_values(d):
    return [
        d.get('finance_request_number', ''),
        d.get('invoice_number', ''),
        d.get('buyer_name', ''),
        d.get('supplier_name', ''),
        d.get('financing_amount', 0.0),
        d.get('earmark_forecast', 0.0),
        d.get('credit_utilization', 0.0),
        d.get('to_be_settled_on_db') if d.get('to_be_settled_on_db') is not None else 'N/A',
        d.get('total_os', 0.0),
        d.get('status_display', ''),
        d.get('finance_status_display', ''),
        d.get('settlement_status', ''),
        d.get('due_date', ''),
        d.get('actual_funding_date', ''),
        d.get('batch_number', ''),
    ]


def _deal_total_row_values(total_row):
    return [
        'Total', '', '', '',
        total_row.get('financing_amount', 0.0),
        total_row.get('earmark_forecast', 0.0),
        total_row.get('credit_utilization', 0.0),
        total_row.get('to_be_settled_on_db', 0.0),
        total_row.get('total_os', 0.0),
        '', '', '', '', '', '',
    ]


def build_deal_detail_workbook(deal_rows, total_row):
    """构建单个买卖方配对的 Deal 明细 Excel，末尾追加合计行。"""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Deal Detail'
    _write_header(ws, _DEAL_HEADERS)
    for d in deal_rows:
        _write_data_row(ws, _deal_row_values(d))
    _write_data_row(ws, _deal_total_row_values(total_row))
    _autofit_columns(ws)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
