"""FCB 服务层 — 客户 CRUD + 交易数据导出"""
import io
import json
import zipfile
import requests
import pandas as pd
from datetime import datetime
from bson import ObjectId
from flask import current_app

import logging

logger = logging.getLogger(__name__)


# n8n API 配置
N8N_INVOICE_API_URL = 'https://n8n.air8.cn/webhook/exportFcbRefactoringInvoices'
N8N_INVOICE_API_TOKEN = 'd5173b84678c11f0995006a0ea58daac'

# Excel 导出列顺序（24列）
EXCEL_COLUMNS = [
    'Client Number',
    'Client Customer #',
    'Customer Name',
    'Customer Address1',
    'Customer Address2',
    'Customer City',
    'Customer State',
    'Customer Zip',
    'Customer Phone',
    'Invoice Number',
    'Invoice Amount',
    'Invoice Date',
    'Client Terms Code',
    'Client Terms Description',
    'Merchandise Amount',
    'Customer Store #',
    'Customer PO Number',
    'Customer Department Number',
    'Discount Amount',
    'Original Invoice Number',
    'Risk Code',
    'Discount Code',
    'Invoice As Of Date',
    'Tradestyle',
]

# MongoDB 集合名
COLLECTION_NAME = 'fcb_clients'


def _get_mongo():
    """获取已初始化的 mongo 对象"""
    try:
        from backend.app.extensions import mongo as ext_mongo
        if ext_mongo is not None:
            return ext_mongo
    except Exception:
        pass
    try:
        if hasattr(current_app, 'extensions'):
            mongo = current_app.extensions.get('mongo')
            if mongo is not None:
                return mongo
    except Exception:
        pass
    return None


def _get_collection():
    """获取 fcb_clients 集合"""
    mongo = _get_mongo()
    if mongo is None:
        raise RuntimeError('MongoDB not initialized')
    return mongo[COLLECTION_NAME]


def _serialize_client(doc):
    """将 MongoDB 文档转为可 JSON 序列化的 dict"""
    if doc is None:
        return None
    doc['_id'] = str(doc['_id'])
    if 'created_at' in doc and isinstance(doc['created_at'], datetime):
        doc['created_at'] = doc['created_at'].isoformat()
    if 'updated_at' in doc and isinstance(doc['updated_at'], datetime):
        doc['updated_at'] = doc['updated_at'].isoformat()
    return doc


def _duplicate_query(supplier_code, buyer_air8_code, customer_name, exclude_id=None):
    """构造查重条件：有 buyer_air8_code 用 (supplier_code, buyer_air8_code)，否则用 (supplier_code, customer_name)"""
    if buyer_air8_code:
        query = {'supplier_code': supplier_code, 'buyer_air8_code': buyer_air8_code}
    else:
        query = {'supplier_code': supplier_code, 'customer_name': customer_name}
    if exclude_id:
        query['_id'] = {'$ne': exclude_id}
    return query


def _merge_row(client, txn):
    """合并客户字段 + 交易字段为一行 Excel 数据（24列）"""
    return {
        'Client Number':              client.get('client_number', ''),
        'Client Customer #':          client.get('client_customer_no', ''),
        'Customer Name':              client.get('customer_name', ''),
        'Customer Address1':          client.get('customer_address1', ''),
        'Customer Address2':          client.get('customer_address2', ''),
        'Customer City':              client.get('customer_city', ''),
        'Customer State':             client.get('customer_state', ''),
        'Customer Zip':               client.get('customer_zip', ''),
        'Customer Phone':             client.get('customer_phone', ''),
        'Invoice Number':             txn.get('invoice_number', ''),
        'Invoice Amount':             txn.get('invoice_amount', ''),
        'Invoice Date':               txn.get('invoice_date', ''),
        'Client Terms Code':          client.get('client_terms_code', ''),
        'Client Terms Description':   client.get('client_terms_desc', ''),
        'Merchandise Amount':         txn.get('merchandise_amount', ''),
        'Customer Store #':           client.get('customer_store_no', ''),
        'Customer PO Number':         txn.get('customer_po_number', ''),
        'Customer Department Number': client.get('customer_dept_no', ''),
        'Discount Amount':            txn.get('discount_amount', ''),
        'Original Invoice Number':    txn.get('original_invoice_number', ''),
        'Risk Code':                  txn.get('risk_code', ''),
        'Discount Code':              txn.get('discount_code', ''),
        'Invoice As Of Date':         txn.get('invoice_as_of_date', ''),
        'Tradestyle':                 txn.get('tradestyle', ''),
    }


def match_client_for_transaction(buyer_code, buyer_name, candidates):
    """在同一 supplier_code 下的候选客户记录中，找到与交易匹配的唯一客户。

    候选记录若填了 buyer_air8_code，仅按其与 buyer_code 精确比较判定；
    未填 buyer_air8_code 的记录，退化为 customer_name 与 buyer_name 的规整化比较。

    返回 (matched_client_or_None, is_ambiguous)
    """
    buyer_code = (buyer_code or '').strip()
    buyer_name_norm = (buyer_name or '').strip().lower()

    code_matches = []
    name_matches = []
    for client in candidates:
        client_buyer_code = (client.get('buyer_air8_code') or '').strip()
        if client_buyer_code:
            if buyer_code and client_buyer_code == buyer_code:
                code_matches.append(client)
        else:
            client_name_norm = (client.get('customer_name') or '').strip().lower()
            if client_name_norm and client_name_norm == buyer_name_norm:
                name_matches.append(client)

    # 精确 buyer_air8_code 匹配的证据强于姓名规整化匹配，两者冲突时优先取 code 匹配
    matches = code_matches if code_matches else name_matches

    if len(matches) == 1:
        return matches[0], False
    if len(matches) == 0:
        return None, False
    return None, True


def _set_excel_styles(worksheet):
    """设置 Excel 样式（复用 export_service 的样式模式）"""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    header_font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    # 样式化表头
    for cell in worksheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # 样式化数据行
    data_alignment = Alignment(horizontal='left', vertical='center')
    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        for cell in row:
            cell.border = thin_border
            cell.alignment = data_alignment

    # 自动列宽
    for col in worksheet.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            val = str(cell.value) if cell.value else ''
            max_length = max(max_length, len(val))
        worksheet.column_dimensions[col_letter].width = min(max_length + 4, 40)


def _build_excel(client, transactions):
    """为单个客户生成 Excel 文件，返回 BytesIO"""
    rows = [_merge_row(client, txn) for txn in transactions]
    df = pd.DataFrame(rows, columns=EXCEL_COLUMNS)

    # 将 None / NaN 替换为空字符串
    df = df.fillna('')
    # Invoice Number 去除连字符
    df['Invoice Number'] = df['Invoice Number'].str.replace('-', '', regex=False)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Sheet1')
        _set_excel_styles(writer.sheets['Sheet1'])
    output.seek(0)
    return output


class FCBService:
    """FCB 业务逻辑服务"""

    # ── 客户 CRUD ──────────────────────────────

    @staticmethod
    def list_clients(page=1, page_size=20, keyword=None):
        col = _get_collection()
        query = {}
        if keyword:
            query = {
                '$or': [
                    {'supplier_code': {'$regex': keyword, '$options': 'i'}},
                    {'customer_name': {'$regex': keyword, '$options': 'i'}},
                    {'client_number': {'$regex': keyword, '$options': 'i'}},
                ]
            }
        total = col.count_documents(query)
        skip = (page - 1) * page_size
        cursor = col.find(query).sort('created_at', -1).skip(skip).limit(page_size)
        items = [_serialize_client(doc) for doc in cursor]
        return {
            'list': items,
            'total': total,
            'page': page,
            'pageSize': page_size,
        }

    @staticmethod
    def get_client(client_id):
        col = _get_collection()
        try:
            doc = col.find_one({'_id': ObjectId(client_id)})
        except Exception:
            return None
        return _serialize_client(doc)

    @staticmethod
    def create_client(data):
        col = _get_collection()
        supplier_code = data.get('supplier_code', '').strip()
        if not supplier_code:
            raise ValueError('supplier_code is required')

        buyer_air8_code = data.get('buyer_air8_code', '').strip()
        customer_name = data.get('customer_name', '')

        # 检查唯一性：有 buyer_air8_code 按其查重，否则按 customer_name 查重
        if col.find_one(_duplicate_query(supplier_code, buyer_air8_code, customer_name)):
            raise ValueError('Client with the same supplier_code and buyer already exists')

        now = datetime.utcnow()
        doc = {
            'supplier_code': supplier_code,
            'buyer_air8_code': buyer_air8_code,
            'client_number': data.get('client_number', ''),
            'client_customer_no': data.get('client_customer_no', ''),
            'customer_name': data.get('customer_name', ''),
            'customer_address1': data.get('customer_address1', ''),
            'customer_address2': data.get('customer_address2', ''),
            'customer_city': data.get('customer_city', ''),
            'customer_state': data.get('customer_state', ''),
            'customer_zip': data.get('customer_zip', ''),
            'customer_phone': data.get('customer_phone', ''),
            'client_terms_code': data.get('client_terms_code', ''),
            'client_terms_desc': data.get('client_terms_desc', ''),
            'customer_store_no': data.get('customer_store_no', ''),
            'customer_dept_no': data.get('customer_dept_no', ''),
            'created_at': now,
            'updated_at': now,
        }
        result = col.insert_one(doc)
        doc['_id'] = result.inserted_id
        return _serialize_client(doc)

    @staticmethod
    def update_client(client_id, data):
        col = _get_collection()
        try:
            oid = ObjectId(client_id)
        except Exception:
            return None

        existing = col.find_one({'_id': oid})
        if not existing:
            return None

        # 不允许更新 _id
        data.pop('_id', None)

        supplier_code = data.get('supplier_code', existing.get('supplier_code', '')).strip()
        buyer_air8_code = data.get('buyer_air8_code', existing.get('buyer_air8_code', '')).strip()
        customer_name = data.get('customer_name', existing.get('customer_name', ''))

        if col.find_one(_duplicate_query(supplier_code, buyer_air8_code, customer_name, exclude_id=oid)):
            raise ValueError('Client with the same supplier_code and buyer already exists')

        # 将规整化后的值写回 data，避免存入未 strip 的原始值（仅限调用方实际提交的字段）
        if 'supplier_code' in data:
            data['supplier_code'] = supplier_code
        if 'buyer_air8_code' in data:
            data['buyer_air8_code'] = buyer_air8_code

        data['updated_at'] = datetime.utcnow()

        result = col.find_one_and_update(
            {'_id': oid},
            {'$set': data},
            return_document=True,
        )
        return _serialize_client(result)

    @staticmethod
    def delete_client(client_id):
        col = _get_collection()
        try:
            oid = ObjectId(client_id)
        except Exception:
            return False
        result = col.delete_one({'_id': oid})
        return result.deleted_count > 0

    # ── 导出 ──────────────────────────────────

    @staticmethod
    def export_transactions(financing_nos=None):
        """导出交易数据为 Excel（单客户）或 ZIP（多客户）"""

        # 1. 调用 n8n API
        transactions = FCBService._call_n8n_api(financing_nos)
        if not transactions:
            raise ValueError('No transaction data found')

        logger.info(
            'FCB export: n8n returned %d transaction(s); sample keys: %s; sample row: %s',
            len(transactions),
            list(transactions[0].keys()),
            transactions[0],
        )

        # 2. 按 supplier_code 分组
        groups = {}
        for txn in transactions:
            code = txn.get('supplier_code', '')
            groups.setdefault(code, []).append(txn)

        # 3. 批量查询候选客户（一个 supplier_code 可能对应多个 buyer）
        col = _get_collection()
        supplier_codes = list(groups.keys())
        candidates_map = {}
        for c in col.find({'supplier_code': {'$in': supplier_codes}}):
            candidates_map.setdefault(c['supplier_code'], []).append(c)

        # 4. 逐笔交易匹配客户，按命中的客户 _id 重新分组
        matched_groups = {}
        unmatched = []
        for code, txns in groups.items():
            candidates = candidates_map.get(code, [])
            for txn in txns:
                client, ambiguous = match_client_for_transaction(
                    txn.get('buyer_code', ''), txn.get('buyer_name', ''), candidates
                )
                if not client:
                    unmatched.append({
                        'supplier_code': code,
                        'buyer_code': txn.get('buyer_code', ''),
                        'buyer_name': txn.get('buyer_name', ''),
                        'ambiguous': ambiguous,
                    })
                    logger.warning(
                        'FCB export unmatched: supplier_code=%s txn_keys=%s txn_buyer_code=%r txn_buyer_name=%r '
                        'candidates=%s',
                        code,
                        list(txn.keys()),
                        txn.get('buyer_code', ''),
                        txn.get('buyer_name', ''),
                        [
                            {
                                '_id': str(c.get('_id')),
                                'buyer_air8_code': c.get('buyer_air8_code', ''),
                                'customer_name': c.get('customer_name', ''),
                            }
                            for c in candidates
                        ],
                    )
                    continue
                cid = client['_id']
                matched_groups.setdefault(cid, {'client': client, 'txns': []})
                matched_groups[cid]['txns'].append(txn)

        if not matched_groups:
            detail = ', '.join(
                f"{u['supplier_code']}/{u['buyer_code'] or u['buyer_name']}" for u in unmatched
            )
            raise ValueError(f'No client records matched for: {detail}')

        if unmatched:
            logger.warning('FCB export: %d transaction(s) unmatched: %s', len(unmatched), unmatched)

        # 5. 生成 Excel 文件
        files = []
        today = datetime.utcnow().strftime('%Y%m%d')
        for group in matched_groups.values():
            client = group['client']
            excel_buf = _build_excel(client, group['txns'])
            client_number = client.get('client_number', 'unknown')
            customer_name = client.get('customer_name', 'unknown')
            # 清理文件名中的非法字符
            safe_name = ''.join(c if c.isalnum() or c in ' _-' else '_' for c in customer_name)
            filename = f'{client_number}_{safe_name}_{today}.xlsx'
            files.append((filename, excel_buf))

        # 6. 返回
        if len(files) == 1:
            fname, buf = files[0]
            return (
                buf,
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                fname,
            )
        else:
            zip_buf = io.BytesIO()
            with zipfile.ZipFile(zip_buf, 'w', zipfile.ZIP_DEFLATED) as zf:
                for fname, buf in files:
                    zf.writestr(fname, buf.read())
            zip_buf.seek(0)
            return (
                zip_buf,
                'application/zip',
                f'fcb_export_{today}.zip',
            )

    @staticmethod
    def _call_n8n_api(financing_nos):
        """调用 n8n API 获取交易数据

        API: GET /webhook/exportFcbRefactoringInvoices?token=...&financingNos=FN1,FN2
        响应格式: [{"data":"[{json string}]"}]
        """
        if not financing_nos:
            raise ValueError('Please provide financing numbers')

        financing_nos_str = ','.join(financing_nos)

        try:
            response = requests.get(
                N8N_INVOICE_API_URL,
                params={
                    'token': N8N_INVOICE_API_TOKEN,
                    'financingNos': financing_nos_str,
                },
                timeout=120,
                verify=False,
            )
            response.raise_for_status()
            result = response.json()

            # 解析嵌套 JSON: [{"data":"[{...}]"}]
            if isinstance(result, list) and result:
                data_str = result[0].get('data', '[]')
                if isinstance(data_str, str):
                    return json.loads(data_str)
                if isinstance(data_str, list):
                    return data_str
            return []
        except requests.RequestException as e:
            raise ValueError(f'Failed to call n8n API: {str(e)}')
